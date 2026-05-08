import os
import io  # <-- Nueva librería para manejar memoria
from flask import Flask, render_template, request, redirect, url_for, send_file
from flask_sqlalchemy import SQLAlchemy
from datetime import datetime
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
import pandas as pd

app = Flask(__name__)

# 1. Configuración de la Base de Datos
app.config['SQLALCHEMY_DATABASE_URI'] = 'sqlite:///ast.db'
app.config['SQLALCHEMY_TRACK_MODIFICATIONS'] = False

db = SQLAlchemy(app)

# 2. Modelo de Datos
class RegistroAST(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    responsable = db.Column(db.String(100), nullable=False)
    area = db.Column(db.String(100), nullable=False)
    tarea = db.Column(db.Text, nullable=False)
    riesgo = db.Column(db.String(50), nullable=False)
    control = db.Column(db.String(100), nullable=False)
    fecha = db.Column(db.DateTime, default=datetime.now)

with app.app_context():
    db.create_all()

# --- RUTAS DE LA APLICACIÓN ---

@app.route('/')
def home():
    return render_template('index.html')

@app.route('/ast', methods=['POST'])
def crear_ast():
    nuevo_ast = RegistroAST(
        responsable = request.form.get('responsable'),
        area = request.form.get('area'),
        tarea = request.form.get('tarea'),
        riesgo = request.form.get('riesgo'),
        control = request.form.get('control')
    )
    db.session.add(nuevo_ast)
    db.session.commit()
    return redirect(url_for('ver_registros'))

@app.route('/registros')
def ver_registros():
    todos_los_ast = RegistroAST.query.order_by(RegistroAST.fecha.desc()).all()
    return render_template('ver_registros.html', registros=todos_los_ast)

@app.route('/eliminar/<int:id>')
def eliminar_ast(id):
    registro_a_borrar = RegistroAST.query.get_or_404(id)
    db.session.delete(registro_a_borrar)
    db.session.commit()
    return redirect(url_for('ver_registros'))

@app.route('/exportar')
def exportar_excel():
    registros = RegistroAST.query.all()
    datos = []
    for r in registros:
        datos.append({
            "FECHA Y HORA": r.fecha.strftime("%d/%m/%Y %H:%M"),
            "ID": r.id,
            "RESPONSABLE": r.responsable,
            "ÁREA": r.area,
            "TAREA": r.tarea,
            "RIESGO": r.riesgo,
            "MEDIDA DE CONTROL": r.control
        })
    
    df = pd.DataFrame(datos)
    
    # --- CAMBIO CLAVE PARA RENDER ---
    output = io.BytesIO() # Creamos un archivo "virtual" en la memoria RAM
    
    # Usamos 'output' en lugar del nombre del archivo
    with pd.ExcelWriter(output, engine='openpyxl') as writer:
        df.to_excel(writer, index=False, sheet_name='Registros_Seguridad')
        
        ws = writer.sheets['Registros_Seguridad']
        
        # Estilos (Tu configuración profesional de US Technologies)
        header_fill = PatternFill(start_color="2C3E50", end_color="2C3E50", fill_type="solid")
        header_font = Font(color="FFFFFF", bold=True, size=12)
        center_alignment = Alignment(horizontal="center", vertical="center")
        border = Border(left=Side(style='thin'), right=Side(style='thin'), 
                        top=Side(style='thin'), bottom=Side(style='thin'))

        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = center_alignment
            cell.border = border

        for column in ws.columns:
            max_length = 0
            column_letter = column[0].column_letter
            for cell in column:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except: pass
            ws.column_dimensions[column_letter].width = (max_length + 4)

        ws.auto_filter.ref = ws.dimensions
        ws.freeze_panes = "A2"

    output.seek(0) # Volvemos al inicio del archivo virtual para poder enviarlo
    
    nombre_descarga = f"Reporte_AST_{datetime.now().strftime('%d-%m-%Y')}.xlsx"
    
    return send_file(
        output, 
        as_attachment=True, 
        download_name=nombre_descarga,
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )

if __name__ == '__main__':
    port = int(os.environ.get('PORT', 5000))
    app.run(host='0.0.0.0', port=port)
